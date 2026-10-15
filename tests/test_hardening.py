"""Regression tests for the production-hardening work: ledger, financial calculations against
hand-computed values, new maker-checker flows, security controls and audit integrity."""
import io
import os
import re
import sys
from datetime import datetime, timedelta
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from harness import (ADMIN, OWNER_BASE, DBNAME, M, PW, app, approve_latest, balance, cleanup, client, db,  # noqa: E402
                     deposit, form_token, idem, invite_employee, post, q, setup_bank, upload)

results = []


def check(code, title, ok, detail=""):
    results.append((code, ok, title, detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {code} {title} :: {detail}")


def D(x):
    return Decimal(str(x))


try:
    S, cid, aid, acc_no = setup_bank()
    off, sen, mgr, admin = S["officer"], S["senior"], S["manager"], S["admin"]
    roles = q(lambda: {r.name: r.id for r in M.Role.query.all()})
    b1 = q(lambda: M.Branch.query.first().id)

    # ---------------- money parsing & formulas (unit level) --------------------------
    from bank.money import allocate_payment, emi, fd_maturity, parse_money, q2, simple_interest
    check("U1", "EMI ₹1,00,000 @ 12% × 12 = ₹8,884.88", emi(100000, 12, 12) == D("8884.88"), emi(100000, 12, 12))
    check("U2", "EMI ₹2,00,000 @ 9.5% × 12 = ₹17,536.70", emi(200000, D("9.5"), 12) == D("17536.70"), emi(200000, D("9.5"), 12))
    check("U3", "EMI at 0.0001% ≈ principal / n", emi(120000, D("0.0001"), 24) == D("5000.01"), emi(120000, D("0.0001"), 24))
    check("U4", "FD ₹50,000 @ 6.8% × 12m quarterly = ₹53,487.69", fd_maturity(50000, D("6.8"), 12) == D("53487.69"))
    check("U5", "Interest ₹1,00,000 @ 12% × 30 days = ₹986.30", simple_interest(100000, 12, 30) == D("986.30"))
    check("U6", "Half-up rounding: ₹2.675 → ₹2.68 (float gives 2.67)", q2("2.675") == D("2.68") and round(2.675, 2) == 2.67)
    bad = ["nan", "inf", "-1", "0", "0.001", "1e3", "abc", "", "10000000000.01", "₹ 5"]
    check("U7", "Invalid amounts refused", all(parse_money(x) is None for x in bad[:-1]), [x for x in bad[:-1] if parse_money(x) is not None])
    check("U8", "Valid amounts parsed exactly", [parse_money(x) for x in ("1", "100", "999.99", "1,00,000")] ==
          [D("1.00"), D("100.00"), D("999.99"), D("100000.00")])
    check("U9", "Payment allocation penalty → interest → principal",
          allocate_payment(1000, 50, 200, 5000) == (D("50.00"), D("200.00"), D("750.00"), D("0.00")))

    # ---------------- ledger: every posting balanced; product balances reconcile -------
    deposit(off, acc_no, "999.99")
    deposit(off, acc_no, "1")
    deposit(off, acc_no, "100", "Debit")
    deposit(off, acc_no, "1000000")  # high value → approval, not posted yet
    n_large = q(lambda: M.ApprovalRequest.query.filter_by(request_type="large_txn").count())
    check("L0", "₹10,00,000 deposit routed to maker-checker", n_large == 1, f"{n_large} request(s)")
    approve_latest(S)
    check("L1", "Balance exact after ₹999.99 + ₹1 − ₹100 + ₹10,00,000", balance(aid) == D("1000900.99"), balance(aid))

    # ---------------- loan: disbursal, penal charge, allocation, closure --------------
    post(off, "/loans/new", {"customer": cid, "loan_type": "personal", "account_id": aid, "amount": "100000",
                             "interest_rate": "12", "tenure_months": "12"}, f"/loans/new?customer={cid}")
    approve_latest(S)
    lid = q(lambda: M.Loan.query.filter_by(loan_type="personal").first().id)
    with app.app_context():
        loan = db.session.get(M.Loan, lid)
        check("LN1", "Disbursed: outstanding = sanctioned = disbursed", loan.outstanding == loan.amount == loan.disbursed_amount == D("100000.00"))
        # make the first EMI 10 days overdue and the last payment 40 days ago
        loan.disbursed_on = loan.disbursed_on - timedelta(days=40)
        loan.next_due_date = loan.disbursed_on + timedelta(days=30)
        db.session.commit()
        from bank.services import loan_dues
        from bank.utils import local_today
        dues = loan_dues(loan, local_today())
        exp_int = simple_interest(D("100000"), 12, 40)
        exp_pen = simple_interest(loan.emi, 24, 10)
        check("LN2", "Interest for 40 days = ₹1,315.07", dues["interest"] == exp_int == D("1315.07"), dues["interest"])
        check("LN3", "Penal charge: EMI × 24% × 10 / 365", dues["penalty"] == exp_pen, f"{dues['penalty']} expected {exp_pen}")
        emi_amt = loan.emi
    post(sen, f"/loans/{lid}/repay", {"amount": str(emi_amt + exp_pen), "mode": "Cash"}, f"/loans/{lid}")
    with app.app_context():
        r = M.LoanRepayment.query.filter_by(loan_id=lid).first()
        ok = r and r.penalty == exp_pen and r.interest_component == exp_int and r.principal_component == emi_amt - exp_int
        check("LN4", "Repayment allocated penalty → interest → principal", bool(ok),
              f"pen {r.penalty if r else None} int {r.interest_component if r else None} prin {r.principal_component if r else None}")
        out_before = db.session.get(M.Loan, lid).outstanding
    post(sen, f"/loans/{lid}/repay", {"amount": str(emi_amt + exp_pen), "mode": "Cash"}, f"/loans/{lid}")  # same day again
    with app.app_context():
        r2 = M.LoanRepayment.query.filter_by(loan_id=lid).order_by(M.LoanRepayment.id.desc()).first()
        check("LN5", "Second same-day payment: no interest, no repeat penal charge", r2.interest_component == 0 and r2.penalty == 0,
              f"int {r2.interest_component} pen {r2.penalty}")
        dues = loan_dues(db.session.get(M.Loan, lid), local_today())
    post(sen, f"/loans/{lid}/repay", {"amount": str(dues["payoff"]), "mode": "Cash"}, f"/loans/{lid}")
    with app.app_context():
        loan = db.session.get(M.Loan, lid)
        check("LN6", "Exact payoff closes the loan at zero", loan.status == "Closed" and loan.outstanding == 0,
              f"{loan.status} {loan.outstanding}")
        paid = sum(r.amount for r in loan.repayments)
        parts = sum(r.principal_component + r.interest_component + r.penalty for r in loan.repayments)
        check("LN7", "Σ repayments = Σ (principal + interest + penal)", paid == parts, f"{paid} vs {parts}")
        prin = sum(r.principal_component for r in loan.repayments)
        check("LN8", "Σ principal repaid = sanctioned amount", prin == loan.amount, f"{prin} vs {loan.amount}")

    # ---------------- interest-rate change via maker-checker --------------------------
    post(off, "/loans/new", {"customer": cid, "loan_type": "home", "account_id": aid, "amount": "500000",
                             "interest_rate": "9", "tenure_months": "120"}, f"/loans/new?customer={cid}")
    approve_latest(S)
    hid = q(lambda: M.Loan.query.filter_by(loan_type="home").first().id)
    post(sen, f"/loans/{hid}/rate", {"new_rate": "10.5", "reason": "Repo rate revision"}, f"/loans/{hid}")
    before = q(lambda: (db.session.get(M.Loan, hid).interest_rate, db.session.get(M.Loan, hid).emi))
    check("RC1", "Rate unchanged until approval", before[0] == D("9"), before)
    rid = q(lambda: M.ApprovalRequest.query.filter_by(request_type="loan_rate_change").first().id)
    post(S["manager2"], f"/approvals/{rid}", {"action": "review", "remarks": "ok"})
    post(mgr, f"/approvals/{rid}", {"action": "approve", "remarks": "ok"})
    after = q(lambda: (db.session.get(M.Loan, hid).interest_rate, db.session.get(M.Loan, hid).emi))
    check("RC2", "Approved: rate 10.5% and EMI recomputed on 120 months", after[0] == D("10.5") and after[1] == emi(500000, D("10.5"), 120),
          after)

    # ---------------- FD: premature closure quote and payout --------------------------
    post(off, "/deposits/new", {"customer": cid, "account_id": aid, "principal": "100000", "tenure_months": "12"},
         f"/deposits/new?customer={cid}")
    approve_latest(S)
    with app.app_context():
        fd = M.FixedDeposit.query.filter_by(status="Active").first()
        fd.start_date = fd.start_date - timedelta(days=183)
        db.session.commit()
        fid = fd.id
        from bank.services import premature_quote
        quote = premature_quote(fd)
        months = Decimal(183) * 12 / Decimal(365)
        expected = fd_maturity(D("100000"), D("5.8"), months)
        check("FD1", "Early-closure quote at 6.8% − 1% for 183 days", quote["rate"] == D("5.80") and quote["value"] == expected,
              f"{quote['value']} expected {expected}")
    bal_before = balance(aid)
    post(admin, f"/deposits/{fid}/close", {"reason": "Medical emergency"}, f"/deposits/{fid}/close")
    approve_latest(S)
    with app.app_context():
        fd = db.session.get(M.FixedDeposit, fid)
        check("FD2", "Early closure paid the quoted value into the account",
              fd.status == "Closed" and fd.premature and balance(aid) - bal_before == fd.paid_out_amount == expected,
              f"{fd.status} paid {fd.paid_out_amount}")

    # ---------------- savings interest from transaction history -----------------------
    with app.app_context():
        from bank.services import post_transaction, savings_interest_preview
        from bank.utils import local_today
        cust = db.session.get(M.Customer, cid)
        acc2 = M.Account(account_number="999000000001", customer_id=cid, branch_id=b1, account_type="Savings",
                         balance=0, interest_rate=D("3"), status="Active")
        db.session.add(acc2)
        db.session.flush()
        today = local_today()
        d0, d1 = today - timedelta(days=100), today - timedelta(days=10)
        noon = lambda d: datetime.combine(d, datetime.min.time()) + timedelta(hours=6)  # noqa: E731  (06:00 UTC = 11:30 IST)
        post_transaction(acc2, "Credit", D("100000"), "Cash", "Deposit", "opening", when=noon(today - timedelta(days=120)))
        post_transaction(acc2, "Credit", D("50000"), "Cash", "Deposit", "top-up", when=noon(today - timedelta(days=50)))
        db.session.commit()
        rows = {r["account_number"]: r for r in savings_interest_preview(d0, d1, [b1])}
        days = (d1 - d0).days + 1
        product = D(100000) * days + D(50000) * ((d1 - (today - timedelta(days=50))).days + 1)
        exp = q2(product * 3 / 100 / 365)
        got = rows.get("999000000001", {}).get("interest")
        check("SI1", "Savings interest = Σ daily balance × 3% / 365", got == exp, f"{got} expected {exp}")
        start_s, end_s = d0.isoformat(), d1.isoformat()
    acct_client = invite_employee(admin, "Accountant One", "accountant@bank.test", roles["Accountant"], b1)
    post(acct_client, "/accounting/interest", {"start": start_s, "end": end_s}, f"/accounting/interest?start={start_s}&end={end_s}")
    irid = q(lambda: M.ApprovalRequest.query.filter_by(request_type="interest_posting").first().id)
    post(mgr, f"/approvals/{irid}", {"action": "review", "remarks": "checked"})
    r = post(mgr, f"/approvals/{irid}", {"action": "approve", "remarks": "x"})
    check("SI2", "Branch Manager cannot give final approval on interest (needs level 90)",
          q(lambda: db.session.get(M.ApprovalRequest, irid).status) == "Reviewed")
    post(admin, f"/approvals/{irid}", {"action": "approve", "remarks": "approved"})
    with app.app_context():
        a2 = M.Account.query.filter_by(account_number="999000000001").one()
        check("SI3", "Interest credited once approved", a2.balance == D("150000") + exp, a2.balance)
    post(acct_client, "/accounting/interest", {"start": start_s, "end": end_s}, f"/accounting/interest?start={start_s}&end={end_s}")
    n_runs = q(lambda: M.ApprovalRequest.query.filter_by(request_type="interest_posting").count())
    check("SI4", "Same period cannot be submitted again", n_runs == 1, f"{n_runs} requests")

    # ---------------- reconciliation after everything ---------------------------------
    with app.app_context():
        from bank.ledger import reconciliation
        checks = reconciliation()
        for c in checks:
            check("REC", c["name"], c["ok"], f"ledger {c['ledger']} / books {c['books']}")
        unbalanced = db.session.execute(db.text(
            "select count(*) from (select entry_id from journal_line group by entry_id having sum(debit) <> sum(credit)) x")).scalar()
        check("REC", "Every journal entry balances", unbalanced == 0, f"{unbalanced} unbalanced")

    # ---------------- account freeze via maker-checker --------------------------------
    post(sen, f"/accounts/{aid}/freeze", {"reason": "Court order"}, f"/customers/{cid}?tab=accounts")
    approve_latest(S)  # senior made it → manager2 / manager decide
    rid = q(lambda: M.ApprovalRequest.query.filter_by(request_type="account_freeze").first())
    if rid.status != "Approved":
        post(S["manager2"], f"/approvals/{rid.id}", {"action": "review", "remarks": "ok"})
        post(mgr, f"/approvals/{rid.id}", {"action": "approve", "remarks": "ok"})
    st = q(lambda: db.session.get(M.Account, aid).status)
    check("FR1", "Account frozen after approval", st == "Frozen", st)
    b0 = balance(aid)
    deposit(off, acc_no, "500", "Debit")
    deposit(off, acc_no, "700")
    check("FR2", "Frozen: debit blocked, credit allowed", balance(aid) == b0 + 700, f"{b0} -> {balance(aid)}")

    # ---------------- customer details change via maker-checker -----------------------
    post(sen, f"/customers/{cid}/edit", {"mobile": "9123456789", "email": "new@example.com", "address": "1 New St",
                                         "city": "Chennai", "occupation": "Engineer", "risk_category": "Medium"})
    m_before = q(lambda: db.session.get(M.Customer, cid).mobile)
    raw_payload = q(lambda: M.ApprovalRequest.query.filter_by(request_type="customer_update").first().payload_json)
    check("CU1", "Mobile unchanged until approved; request payload encrypted",
          m_before == "9876543210" and raw_payload.startswith("enc:") and "9123456789" not in raw_payload)
    rid = q(lambda: M.ApprovalRequest.query.filter_by(request_type="customer_update").first().id)
    post(S["manager2"], f"/approvals/{rid}", {"action": "review", "remarks": "ok"})
    post(mgr, f"/approvals/{rid}", {"action": "approve", "remarks": "ok"})
    check("CU2", "Mobile updated after approval", q(lambda: db.session.get(M.Customer, cid).mobile) == "9123456789")

    # ---------------- role permission change by a non-super-admin needs approval -----------
    bankadmin = invite_employee(admin, "Bank Admin", "bankadmin@bank.test", roles["Bank Admin"], b1)
    clerk_role = roles["Clerk"]
    before = q(lambda: db.session.get(M.Role, clerk_role).perm("reports").can_view)
    t, html = form_token(bankadmin, f"/users/roles/{clerk_role}")
    form = {"_csrf": t, "_idem": idem(html), "reports:view": "on", "customer_accounts:view": "on", "transactions:view": "on",
            "transactions:create": "on", "kyc:view": "on", "kyc:create": "on", "tickets:view": "on", "tickets:create": "on",
            "alerts:view": "on"}
    bankadmin.post(f"/users/roles/{clerk_role}", data=form)
    after = q(lambda: db.session.get(M.Role, clerk_role).perm("reports").can_view)
    pend = q(lambda: M.ApprovalRequest.query.filter_by(request_type="role_permissions").count())
    check("RP1", "Bank Admin's permission change waits for approval", before == after is False and pend == 1, f"pending {pend}")

    # ---------------- invite links ------------------------------------------------------
    r = post(admin, "/users/new", {"name": "Invitee", "email": "invitee@bank.test", "role_id": roles["Clerk"], "branch_id": b1})
    link = re.search(r'/invite/([A-Za-z0-9_\-]+)', r.get_data(as_text=True)).group(1)
    anon = app.test_client()
    weak = post(anon, f"/invite/{link}", {"password": "short1!", "confirm": "short1!"})
    check("IV1", "Weak password refused on invite", "Password needs" in weak.get_data(as_text=True))
    ok = post(anon, f"/invite/{link}", {"password": PW, "confirm": PW})
    reuse = anon.get(f"/invite/{link}")
    check("IV2", "Invite link works once, then 404", "Your password is set" in ok.get_data(as_text=True) and reuse.status_code == 404,
          reuse.status_code)
    raw_tokens = q(lambda: [t.token_hash for t in M.InviteToken.query.all()])
    check("IV3", "Only hashes of invite tokens are stored", link not in raw_tokens and all(len(h) == 64 for h in raw_tokens))

    # ---------------- MFA code replay -------------------------------------------------
    import pyotp
    mfa_user = invite_employee(admin, "Mfa User", "mfa@bank.test", roles["Clerk"], b1)
    page = mfa_user.get("/settings/mfa").get_data(as_text=True)
    secret = re.search(r'<p class="mono">([A-Z2-7]{16,})</p>', page).group(1)
    t, html = form_token(mfa_user, "/settings/mfa")
    code0 = pyotp.TOTP(secret).now()
    mfa_user.post("/settings/mfa/enable", data={"_csrf": t, "_idem": idem(html), "code": code0})
    c = app.test_client()
    t = form_token(c, "/login")[0]
    c.post("/login", data={"_csrf": t, "email": "mfa@bank.test", "password": PW})
    import time as _time
    code = pyotp.TOTP(secret).at(_time.time() + 30)  # next step: the enrolment code itself is spent
    t = form_token(c, "/login/mfa")[0]
    first = c.post("/login/mfa", data={"_csrf": t, "code": code}, follow_redirects=False)
    c2 = app.test_client()
    t = form_token(c2, "/login")[0]
    c2.post("/login", data={"_csrf": t, "email": "mfa@bank.test", "password": PW})
    t = form_token(c2, "/login/mfa")[0]
    replay = c2.post("/login/mfa", data={"_csrf": t, "code": code}, follow_redirects=False)
    check("MFA1", "Same TOTP code cannot be used twice", first.status_code == 302 and replay.status_code == 200,
          f"first {first.status_code}, replay {replay.status_code}")
    c3 = app.test_client()
    t = form_token(c3, "/login")[0]
    c3.post("/login", data={"_csrf": t, "email": "mfa@bank.test", "password": PW})
    t = form_token(c3, "/login/mfa")[0]
    enrol_replay = c3.post("/login/mfa", data={"_csrf": t, "code": code0}, follow_redirects=False)
    check("MFA2", "The enrolment code cannot be replayed to sign in", enrol_replay.status_code == 200, enrol_replay.status_code)

    # ---------------- documents encrypted at rest, tamper detected ---------------------
    png = b"\x89PNG\r\n\x1a\n" + os.urandom(64)
    upload(admin, cid, "Passport", "P1234567", content=png, filename="passport.png")
    with app.app_context():
        doc = M.Document.query.filter_by(doc_type="Passport").order_by(M.Document.id.desc()).first()
        path = os.path.join(app.config["UPLOAD_FOLDER"], doc.stored_name)
        raw = open(path, "rb").read()
        did = doc.id
    check("DOC1", "Stored file is encrypted (no PNG signature on disk)", not raw.startswith(b"\x89PNG") and png not in raw)
    got = admin.get(f"/kyc/doc/{did}/file")
    check("DOC2", "Download decrypts to the original bytes", got.status_code == 200 and got.data == png)
    bad = upload(admin, cid, "Passport", "X", content=b"MZ\x90\x00 executable", filename="passport.pdf")
    check("DOC3", "Executable renamed .pdf refused", "does not match its type" in bad.get_data(as_text=True))
    js = upload(admin, cid, "Passport", "X", content=b"%PDF-1.4 /JavaScript (app.alert(1))", filename="evil.pdf")
    check("DOC4", "PDF with embedded JavaScript refused", "scripts" in js.get_data(as_text=True))
    with open(path, "r+b") as fh:
        fh.seek(40)
        fh.write(b"XX")
    tampered = admin.get(f"/kyc/doc/{did}/file")
    check("DOC5", "Tampered file is not served", tampered.status_code == 404, tampered.status_code)

    # ---------------- KYC renewal clears expiry ----------------------------------------
    with app.app_context():
        old = M.Document(customer_id=cid, category="KYC", doc_type="Driving Licence", status="Expired",
                         expiry_date=local_today() - timedelta(days=5), uploaded_by=None)
        old.doc_number = "DL1"
        db.session.add(old)
        db.session.get(M.Customer, cid).kyc_status = "Expired"
        db.session.commit()
    post(sen, f"/kyc/doc/{did}/verify", {"decision": "approve"}, f"/kyc/doc/{did}")  # the passport from DOC1
    upload(admin, cid, "Driving Licence", "DL2", content=b"%PDF-1.4 renewed", filename="dl.pdf")
    new_id = q(lambda: M.Document.query.filter_by(doc_type="Driving Licence").order_by(M.Document.id.desc()).first().id)
    post(sen, f"/kyc/doc/{new_id}/verify", {"decision": "approve"}, f"/kyc/doc/{new_id}")
    check("KYC1", "Verified renewal clears the customer's expired KYC", q(lambda: db.session.get(M.Customer, cid).kyc_status) == "Verified",
          q(lambda: db.session.get(M.Customer, cid).kyc_status))

    # ---------------- exports: formula injection, cap, permission ---------------------
    post(admin, "/customers/new", {"name": "=HYPERLINK(\"http://evil\",\"x\")", "pan": "ZZZZZ1111Z", "mobile": "9000000001",
                                   "branch_id": b1, "gender": "Male"})
    x = admin.get("/reports/kyc-status?range=year&export=xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.data))
    vals = [c.value for row in wb.active.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    check("EX1", "Formula-like customer name exported as text", vals and all(v.startswith("'") for v in vals), vals[:1])
    with app.app_context():
        M.Setting.set("max_export_rows", "100")
        db.session.commit()
    big = admin.get("/reports/audit-activity?range=year&export=xlsx", follow_redirects=True)
    check("EX2", "Export above the row cap refused", "limit for bulk extraction" in big.get_data(as_text=True))
    with app.app_context():
        M.Setting.set("max_export_rows", "50000")
        db.session.commit()
    check("EX3", "Clerk cannot export or print statements",
          S["clerk"].get("/reports/loans?export=xlsx").status_code == 403 and S["clerk"].get(f"/customers/{cid}/statement").status_code == 403)

    # ---------------- branch scoping on settings --------------------------------------
    post(admin, "/settings/branches/new", {"code": "BR2", "name": "Branch Two"}, "/settings/")
    b2 = q(lambda: M.Branch.query.filter_by(code="BR2").one().id)
    t, html = form_token(mgr, "/settings/")
    mgr.post("/settings/", data={"_csrf": t, "_idem": idem(html), f"b{b2}_name": "Hijacked", f"b{b2}_dep": "1"})
    check("BR1", "Branch Manager cannot edit another branch", q(lambda: db.session.get(M.Branch, b2).name) == "Branch Two")
    t, html = form_token(mgr, "/settings/")
    sc = mgr.post("/settings/branches/new", data={"_csrf": t, "_idem": idem(html), "code": "BR3", "name": "Rogue"}).status_code
    check("BR2", "Branch Manager cannot open a branch", sc == 403 and not q(lambda: M.Branch.query.filter_by(code="BR3").first()), sc)

    # ---------------- spoofed client IP ------------------------------------------------
    c = app.test_client()
    t = form_token(c, "/login")[0]
    c.post("/login", data={"_csrf": t, "email": "nobody@x.y", "password": "x"}, headers={"X-Forwarded-For": "6.6.6.6"})
    ip = q(lambda: M.LoginHistory.query.order_by(M.LoginHistory.id.desc()).first().ip_address)
    check("IP1", "X-Forwarded-For from an untrusted client is ignored", ip != "6.6.6.6", ip)

    # ---------------- disabled employee & absolute session ----------------------------
    clerk = S["clerk"]
    cid_clerk = q(lambda: M.Employee.query.filter_by(email="clerk@bank.test").one().id)
    post(admin, f"/users/{cid_clerk}/edit", {"action": "save", "role_id": roles["Clerk"], "branch_id": b1, "name": "Clerk"})
    r = clerk.get("/accounts/", follow_redirects=False)
    check("SES1", "Disabled employee's open session is refused", r.status_code in (302, 401), r.status_code)
    old_session = client(ADMIN[0], ADMIN[1])
    with old_session.session_transaction() as sess:
        sess["login_at"] = sess["login_at"] - 13 * 3600
        sess["last_seen"] = datetime.utcnow().timestamp()
    r = old_session.get("/", follow_redirects=False)
    check("SES2", "Session older than 12 h ends even if active", r.status_code == 302 and "/login" in r.headers.get("Location", ""))

    # ---------------- audit log: sealed chain, tamper detected -------------------------
    with app.app_context():
        from bank.audit_chain import seal, verify
        seal()
        v = verify()
        check("AUD1", "Audit hash chain verifies", v["ok"] and v["checked"] > 50, v)
    from sqlalchemy import create_engine, text
    owner = create_engine(f"{OWNER_BASE}/{DBNAME}")
    with owner.begin() as conn:  # an insider with superuser rights bypasses the trigger…
        conn.execute(text("ALTER TABLE audit_log DISABLE TRIGGER audit_log_seal_only"))
        conn.execute(text("UPDATE audit_log SET details = 'edited' WHERE id = (SELECT min(id) + 5 FROM audit_log)"))
        conn.execute(text("ALTER TABLE audit_log ENABLE TRIGGER audit_log_seal_only"))
    owner.dispose()
    with app.app_context():
        v = verify()
        check("AUD2", "…but the edit is detected by the chain", not v["ok"] and v["broken_at"], v)

    # ---------------- health endpoints -------------------------------------------------
    anon = app.test_client()
    check("H1", "Liveness and readiness respond without login",
          anon.get("/healthz").status_code == 200 and anon.get("/readyz").get_json()["status"] == "ready",
          anon.get("/readyz").get_json())
finally:
    cleanup()

failed = [r for r in results if not r[1]]
print(f"\nhardening: {len(results) - len(failed)}/{len(results)} passed")
sys.exit(1 if failed else 0)
