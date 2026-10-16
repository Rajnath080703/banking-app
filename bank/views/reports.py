import io
from datetime import datetime
from decimal import Decimal

from flask import Blueprint, abort, render_template, request, send_file
from flask_login import current_user, login_required
from sqlalchemy import func
from sqlalchemy.orm import joinedload

from ..extensions import db, limiter
from ..money import q2
from ..models import (Account, ApprovalRequest, AuditLog, Branch, Customer, Document, Employee, FeeCharge,
                      FixedDeposit, InterestTds, Loan, LoanRepayment, Ticket, Transaction)
from ..security import audit, require_perm
from ..utils import date_range_from_request, day_bounds, sql_day, sql_month, to_local

bp = Blueprint("reports", __name__, url_prefix="/reports")


def _branches():
    if current_user.sees_all_branches:
        b = request.args.get("branch", type=int)
        return [b] if b else [x.id for x in Branch.query.all()]
    return [current_user.branch_id]


def _money(v):
    return q2(v or 0)


FORMULA_PREFIXES = ("=", "+", "-", "@", chr(9), chr(13))


def safe_cell(v):
    """Neutralise spreadsheet formula injection: text that starts like a formula
    (=, +, -, @) from customer-entered fields is exported as plain text."""
    if isinstance(v, str) and v.startswith(FORMULA_PREFIXES):
        return "'" + v
    return v


def _export_key():
    return f"export:{current_user.get_id().split(':')[0]}" if current_user.is_authenticated else "anon"


def r_daily_txn(start, end, br):
    s, e = day_bounds(start, end)
    day = sql_day(Transaction.timestamp)
    rows = (db.session.query(day, func.count(),
                             func.sum(db.case((Transaction.txn_type == "Credit", Transaction.amount), else_=0)),
                             func.sum(db.case((Transaction.txn_type == "Debit", Transaction.amount), else_=0)))
            .filter(Transaction.branch_id.in_(br), Transaction.timestamp.between(s, e))
            .group_by(day).order_by(day.desc()).all())
    return (["Date", "Transactions", "Credits (₹)", "Debits (₹)", "Net (₹)"],
            [[d, n, _money(c), _money(dd), _money((c or 0) - (dd or 0))] for d, n, c, dd in rows])


def r_monthly_txn(start, end, br):
    s, e = day_bounds(start, end)
    m = sql_month(Transaction.timestamp)
    rows = (db.session.query(m, func.count(),
                             func.sum(db.case((Transaction.txn_type == "Credit", Transaction.amount), else_=0)),
                             func.sum(db.case((Transaction.txn_type == "Debit", Transaction.amount), else_=0)))
            .filter(Transaction.branch_id.in_(br), Transaction.timestamp.between(s, e))
            .group_by(m).order_by(m.desc()).all())
    return (["Month", "Transactions", "Credits (₹)", "Debits (₹)", "Net (₹)"],
            [[d, n, _money(c), _money(dd), _money((c or 0) - (dd or 0))] for d, n, c, dd in rows])


def _txn_list(start, end, br, ttype):
    s, e = day_bounds(start, end)
    rows = (Transaction.query.options(joinedload(Transaction.account), joinedload(Transaction.customer))
            .filter(Transaction.branch_id.in_(br), Transaction.txn_type == ttype,
                                     Transaction.timestamp.between(s, e))
            .order_by(Transaction.timestamp.desc()).limit(5000).all())
    return (["Date", "Reference", "Account", "Customer", "Channel", "Category", "Amount (₹)"],
            [[to_local(t.timestamp).strftime("%Y-%m-%d %H:%M"), t.txn_ref, t.account.account_number, t.customer.name,
              t.channel, t.category, _money(t.amount)] for t in rows])


def r_deposits(start, end, br):
    cols, rows = _txn_list(start, end, br, "Credit")
    fds = (FixedDeposit.query.options(joinedload(FixedDeposit.customer))
           .filter(FixedDeposit.branch_id.in_(br), FixedDeposit.start_date.between(start, end)).all())
    rows += [[fd.start_date.isoformat(), fd.fd_number, "FD", fd.customer.name, "Transfer",
              f"FD {fd.tenure_months}m @ {fd.interest_rate}%", _money(fd.principal)] for fd in fds]
    return cols, sorted(rows, key=lambda r: r[0], reverse=True)


def r_withdrawals(start, end, br):
    return _txn_list(start, end, br, "Debit")


def r_loans(start, end, br):
    rows = Loan.query.options(joinedload(Loan.customer)).filter(Loan.branch_id.in_(br)).order_by(Loan.created_at.desc()).all()
    return (["Loan No", "Customer", "Type", "Amount (₹)", "Outstanding (₹)", "Rate %", "EMI (₹)", "Status", "Disbursed"],
            [[l.loan_number, l.customer.name, l.type_label, _money(l.amount), _money(l.outstanding),
              l.interest_rate, _money(l.emi), l.status, l.disbursed_on.isoformat() if l.disbursed_on else ""] for l in rows])


def r_collections(start, end, br):
    rows = (LoanRepayment.query.options(joinedload(LoanRepayment.loan).joinedload(Loan.customer)).join(Loan).filter(Loan.branch_id.in_(br), LoanRepayment.paid_on.between(start, end))
            .order_by(LoanRepayment.paid_on.desc()).all())
    return (["Paid On", "Loan No", "Customer", "Type", "Amount (₹)", "Principal (₹)", "Interest (₹)", "Mode", "Status"],
            [[r.paid_on.isoformat(), r.loan.loan_number, r.loan.customer.name, r.loan.type_label, _money(r.amount),
              _money(r.principal_component), _money(r.interest_component), r.mode, r.status] for r in rows])


def r_overdue(start, end, br):
    rows = Loan.query.options(joinedload(Loan.customer)).filter(Loan.branch_id.in_(br), Loan.status == "Overdue").order_by(Loan.overdue_days.desc()).all()
    return (["Loan No", "Customer", "Type", "Outstanding (₹)", "EMI (₹)", "Due Date", "Days Overdue", "Bucket"],
            [[l.loan_number, l.customer.name, l.type_label, _money(l.outstanding), _money(l.emi),
              l.next_due_date.isoformat() if l.next_due_date else "", l.overdue_days,
              "0-30" if l.overdue_days <= 30 else "31-60" if l.overdue_days <= 60 else "61-90" if l.overdue_days <= 90 else "90+ (NPA)"]
             for l in rows])


def r_fees(start, end, br):
    s, e = day_bounds(start, end)
    rows = (FeeCharge.query.options(joinedload(FeeCharge.customer), joinedload(FeeCharge.fee_type), joinedload(FeeCharge.account))
            .filter(FeeCharge.branch_id.in_(br), FeeCharge.created_at.between(s, e)).order_by(FeeCharge.id.desc()).all())
    return (["Date", "Charge", "Customer", "Account", "Fee", "Period", "Fee (₹)", "GST (₹)", "Total (₹)", "Status"],
            [[to_local(c.created_at).date().isoformat(), c.charge_ref, c.customer.name, c.account.account_number,
              c.fee_type.name, c.period or "", _money(c.base_amount), _money(c.gst_amount), _money(c.total), c.status]
             for c in rows])


def r_tds(start, end, br):
    s, e = day_bounds(start, end)
    rows = (InterestTds.query.options(joinedload(InterestTds.customer), joinedload(InterestTds.fd))
            .filter(InterestTds.branch_id.in_(br), InterestTds.created_at.between(s, e)).order_by(InterestTds.id.desc()).all())
    return (["Date", "FY", "Customer", "Customer ID", "PAN", "FD No", "Interest paid (₹)", "TDS rate (%)", "TDS (₹)", "Note"],
            [[to_local(r.created_at).date().isoformat(), r.fin_year, r.customer.name, r.customer.customer_code,
              r.customer.pan_masked if r.pan_available else "Not provided", r.fd.fd_number if r.fd else "",
              _money(r.interest), r.rate, _money(r.tds), r.note or ""] for r in rows])


def r_new_accounts(start, end, br):
    rows = Account.query.options(joinedload(Account.customer)).filter(Account.branch_id.in_(br), Account.opened_on.between(start, end)).order_by(Account.opened_on.desc()).all()
    return (["Opened", "Account No", "Customer", "Type", "Status", "Balance (₹)"],
            [[a.opened_on.isoformat(), a.account_number, a.customer.name, a.account_type, a.status, _money(a.balance)] for a in rows])


def r_closed_accounts(start, end, br):
    rows = Account.query.options(joinedload(Account.customer)).filter(Account.branch_id.in_(br), Account.status == "Closed",
                                Account.closed_on.between(start, end)).order_by(Account.closed_on.desc()).all()
    return (["Closed", "Account No", "Customer", "Type", "Opened"],
            [[a.closed_on.isoformat(), a.account_number, a.customer.name, a.account_type, a.opened_on.isoformat()] for a in rows])


def r_kyc(start, end, br):
    docs = (db.session.query(Document.customer_id, func.count(Document.id),
                             func.count(Document.id).filter(Document.status == "Verified"),
                             func.min(Document.expiry_date))
            .filter(Document.category == "KYC").group_by(Document.customer_id).subquery())
    rows = (db.session.query(Customer.customer_code, Customer.name, Customer.kyc_status,
                             docs.c[1], docs.c[2], docs.c[3])
            .outerjoin(docs, docs.c.customer_id == Customer.id)
            .filter(Customer.branch_id.in_(br)).order_by(Customer.kyc_status, Customer.name).all())
    return (["Customer ID", "Name", "KYC Status", "Documents", "Verified", "Next Expiry"],
            [[code, name, st, n or 0, v or 0, nxt.isoformat() if nxt else ""] for code, name, st, n, v, nxt in rows])


def r_tickets(start, end, br):
    s, e = day_bounds(start, end)
    rows = Ticket.query.options(joinedload(Ticket.customer), joinedload(Ticket.assignee)).filter(Ticket.branch_id.in_(br), Ticket.created_at.between(s, e)).order_by(Ticket.created_at.desc()).all()
    return (["Ticket", "Created", "Customer", "Subject", "Category", "Priority", "Status", "Assigned To"],
            [[t.ticket_number, to_local(t.created_at).strftime("%Y-%m-%d"), t.customer.name, t.subject, t.category, t.priority,
              t.status, t.assignee.name if t.assignee else ""] for t in rows])


def r_branch(start, end, br):
    out = []
    for b in Branch.query.filter(Branch.id.in_(br)).all():
        dep = (db.session.query(func.coalesce(func.sum(Account.balance), 0)).filter(Account.branch_id == b.id, Account.status != "Closed").scalar()
               + db.session.query(func.coalesce(func.sum(FixedDeposit.principal), 0)).filter(FixedDeposit.branch_id == b.id, FixedDeposit.status == "Active").scalar())
        loans = db.session.query(func.coalesce(func.sum(Loan.outstanding), 0)).filter(Loan.branch_id == b.id, Loan.status.in_(["Active", "Overdue"])).scalar()
        new_acc = Account.query.filter(Account.branch_id == b.id, Account.opened_on.between(start, end)).count()
        od = Loan.query.filter(Loan.branch_id == b.id, Loan.status == "Overdue").count()
        out.append([b.name, Customer.query.filter_by(branch_id=b.id).count(), _money(dep), _money(b.target_deposits),
                    f"{dep / b.target_deposits * 100:.1f}%" if b.target_deposits else "", _money(loans), _money(b.target_loans),
                    f"{loans / b.target_loans * 100:.1f}%" if b.target_loans else "", new_acc, od])
    return (["Branch", "Customers", "Deposits (₹)", "Deposit Target (₹)", "Dep. Achieved", "Loans (₹)",
             "Loan Target (₹)", "Loan Achieved", "New Accounts", "Overdue Loans"], out)


def r_approvals(start, end, br):
    s, e = day_bounds(start, end)
    rows = ApprovalRequest.query.options(joinedload(ApprovalRequest.maker), joinedload(ApprovalRequest.checker),
                                         joinedload(ApprovalRequest.approver)).filter(ApprovalRequest.branch_id.in_(br), ApprovalRequest.created_at.between(s, e)) \
        .order_by(ApprovalRequest.created_at.desc()).all()
    return (["Request", "Created", "Type", "Summary", "Maker", "Checker", "Approver", "Status"],
            [[r.request_number, to_local(r.created_at).strftime("%Y-%m-%d"), r.type_label, r.summary, r.maker.name,
              r.checker.name if r.checker else "", r.approver.name if r.approver else "", r.status] for r in rows])


def r_audit(start, end, br):
    s, e = day_bounds(start, end)
    q = AuditLog.query.filter(AuditLog.timestamp.between(s, e))
    if not current_user.sees_all_branches:
        q = q.join(Employee, AuditLog.employee_id == Employee.id).filter(Employee.branch_id.in_(br))
    rows = q.order_by(AuditLog.timestamp.desc()).limit(5000).all()
    return (["Timestamp", "User", "Action", "Module", "Reference", "Details", "IP"],
            [[to_local(a.timestamp).strftime("%Y-%m-%d %H:%M:%S"), a.email or "", a.action, a.module or "", a.reference or "",
              a.details or "", a.ip_address or ""] for a in rows])


# key: (title, description, required module, function)
REPORTS = {
    "daily-transactions": ("Daily Transactions", "Credits, debits and counts per day", "transactions", r_daily_txn),
    "monthly-transactions": ("Monthly Transactions", "Month-wise transaction summary", "transactions", r_monthly_txn),
    "deposits": ("Deposits", "Credits and new fixed deposits", "deposits", r_deposits),
    "fees": ("Fees & Charges", "Fees charged, due, waived and skipped, with GST", "transactions", r_fees),
    "tds": ("TDS on FD Interest", "Interest paid and tax deducted, per customer and financial year", "deposits", r_tds),
    "withdrawals": ("Withdrawals", "All debit transactions", "transactions", r_withdrawals),
    "loans": ("Loans", "Loan book with outstanding balances", "loans", r_loans),
    "loan-collections": ("Loan Collections", "EMI and repayment receipts", "loans", r_collections),
    "overdue-loans": ("Overdue Loans", "Overdue loans with ageing buckets", "loans", r_overdue),
    "new-accounts": ("New Accounts", "Accounts opened in the period", "customer_accounts", r_new_accounts),
    "closed-accounts": ("Closed Accounts", "Accounts closed in the period", "customer_accounts", r_closed_accounts),
    "kyc-status": ("KYC Status", "Customer KYC completeness and expiry", "kyc", r_kyc),
    "customer-tickets": ("Customer Tickets", "Service requests and complaints", "tickets", r_tickets),
    "branch-performance": ("Branch Performance", "Targets vs achievement by branch", "branch_management", r_branch),
    "approval-activity": ("Approval Activity", "Maker-checker requests and outcomes", "loan_approval", r_approvals),
    "audit-activity": ("Audit Activity", "Security and user activity trail", "audit_logs", r_audit),
}


@bp.route("/")
@login_required
@require_perm("reports", "view")
def index():
    available = {k: v for k, v in REPORTS.items() if current_user.can(v[2], "view")}
    return render_template("reports/index.html", reports=available)


@bp.route("/<key>")
@login_required
@require_perm("reports", "view")
@limiter.limit("30 per hour", key_func=_export_key, exempt_when=lambda: not request.args.get("export"))
def view(key):
    if key not in REPORTS:
        abort(404)
    title, desc, module, fn = REPORTS[key]
    if not current_user.can(module, "view"):
        abort(403)
    start, end, rkey = date_range_from_request("month")
    cols, rows = fn(start, end, _branches())
    fmt = request.args.get("export")
    if fmt:
        if fmt not in ("xlsx", "pdf"):
            abort(400)
        if not current_user.can("reports", "export") or not current_user.can(module, "view"):
            audit("Report export denied", "reports", key, fmt, result="denied")
            db.session.commit()
            abort(403)
        from ..models import Setting
        cap = Setting.get_int("max_export_rows")
        if len(rows) > cap:
            audit("Report export refused (too large)", "reports", key, f"{len(rows)} rows > {cap}", result="denied")
            db.session.commit()
            from flask import flash, redirect, url_for
            flash(f"This export has {len(rows):,} rows, above the {cap:,}-row limit for bulk extraction. "
                  f"Narrow the date range or branch and try again.", "warning")
            return redirect(url_for("reports.view", key=key, **{k: v for k, v in request.args.items() if k != "export"}))
        audit("Report exported", "reports", key, f"{fmt} {start}..{end} rows={len(rows)}")
        db.session.commit()
        return _export(fmt, key, title, cols, rows, start, end)
    branches = Branch.query.all() if current_user.sees_all_branches else []
    return render_template("reports/view.html", key=key, title=title, desc=desc, cols=cols, rows=rows[:500],
                           total=len(rows), start=start, end=end, rkey=rkey, branches=branches)


def _export(fmt, key, title, cols, rows, start, end):
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    subtitle = f"{start:%d %b %Y} – {end:%d %b %Y} · Generated by {current_user.name} on {datetime.now():%d %b %Y %H:%M}"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([subtitle])
        ws.append([])
        ws.append(cols)
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
            cell.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([safe_cell(v) for v in r])
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:300]]) + 2
            ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(width, 45)
        ws.freeze_panes = "A5"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{key}_{stamp}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                                topMargin=12 * mm, bottomMargin=12 * mm, title=title)
        styles = getSampleStyleSheet()
        cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
        head = cell.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
        clean = lambda v: str(v).replace("₹", "Rs.").replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
        data = [[Paragraph(clean(c), head) for c in cols]]
        data += [[Paragraph(clean(f"{v:,.2f}" if isinstance(v, (float, Decimal)) else v), cell) for v in r] for r in rows[:3000]]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story = [Paragraph(clean(title), styles["Title"]), Paragraph(clean(subtitle), styles["Normal"]),
                 Spacer(1, 6 * mm), table]
        if len(rows) > 3000:
            story.append(Paragraph(f"Showing first 3000 of {len(rows)} rows — use Excel export for the full data.", styles["Italic"]))
        doc.build(story)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{key}_{stamp}.pdf", mimetype="application/pdf")
    abort(400)
